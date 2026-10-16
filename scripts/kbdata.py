"""Shared loading for the n8n Node description files."""
from __future__ import annotations

import os
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

HERE = Path(__file__).resolve().parent
# The two xlsx files are not in the git repo. Set EB_KB_MASTER and EB_AUGUST_MASTER
# to their paths, or keep the default folder layout.
KB = Path(os.environ.get("EB_KB_MASTER") or HERE.parent / "knowledge_base" / "EB_Knowledge_Base_Master.xlsx")
AUGUST = Path(os.environ.get("EB_AUGUST_MASTER")
              or HERE.parent.parent / "data" / "input" / "EuBI_Master_Nodes_and_Facilities_v1.xlsx")
LINK_STATUS = HERE / "link_status.json"
LINK_OVERRIDES = HERE / "link_overrides.csv"
LINK_REPLACEMENTS = HERE / "link_replacements.csv"   # old_url,new_url,checked_by,date,note
EXTRA_NODE_WEBSITES = HERE / "extra_node_websites.csv"   # node_uri,url,checked_by,date
NODE_FILES = HERE / "node_files"

SHEETS = ("Nodes", "Facilities", "Technologies", "Node Technologies", "Facility Technologies",
          "Services", "Contacts", "Data Enrichment", "Technology Guidance", "Node Services",
          "Official Services", "API Access Routes", "Evidence")
# Evidence rows that describe a facility's current capability in structured form.
EVIDENCE_FIELDS = {"matching.instrument_model", "matching.workflow_support", "matching.technique_configuration",
                   "matching.instrument_configuration", "matching.biosafety", "matching.extended_method_capability"}
EVIDENCE_SKIP_ROUTES = {"HISTORY_ONLY", "AUDIT_ONLY", "HELD_OR_UNRESOLVED_CONTEXT", "HELD_MEMBERSHIP_CONTEXT",
                        "DATED_PRICE_CONTEXT_NOT_QUOTE", "UTILITY_EXCLUDED", "LINEAGE_PROVENANCE_ONLY"}


def s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    return str(v).strip()


def rows(wb, sheet: str) -> list[dict]:
    it = wb[sheet].iter_rows(values_only=True)
    head = [s(h) for h in next(it)]
    out = []
    for r in it:
        if not r or all(v in (None, "") for v in r):
            continue
        out.append({h: v for h, v in zip(head, r) if h})
    return out


def load_kb(path: Path = KB) -> dict[str, list[dict]]:
    print(f"Loading {path.name} (about a minute) ...", flush=True)
    wb = openpyxl.load_workbook(path, read_only=True)
    out = {}
    for name in SHEETS:
        if name not in wb.sheetnames:
            continue
        data = rows(wb, name)
        if name == "Evidence":  # 33,000 rows; keep only structured facility capability facts
            data = [r for r in data if r.get("field_or_relation") in EVIDENCE_FIELDS
                    and s(r.get("navigator_route")) not in EVIDENCE_SKIP_ROUTES]
        out[name] = data
    return out


def load_august(path: Path = AUGUST) -> dict[str, list[dict]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return {"Nodes": rows(wb, "Nodes"), "Facilities": rows(wb, "Facilities")}


def norm(v) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", s(v).lower())).strip()


def august_node_map(kb, aug) -> dict[str, dict]:
    """August Node row per KB node_uri, matched through shared EBI-FAC IDs, then by name."""
    fac_node = {f["ebi_fac_id"]: f.get("node_uri") for f in kb["Facilities"] if f.get("ebi_fac_id")}
    votes: dict[str, dict[str, int]] = {}
    for f in aug["Facilities"]:
        uri = fac_node.get(s(f.get("Facility ID")))
        if uri:
            votes.setdefault(s(f.get("Node")), {}).setdefault(uri, 0)
            votes[s(f.get("Node"))][uri] += 1
    by_name = {norm(n.get("node_name")): n["node_uri"] for n in kb["Nodes"] if n.get("node_uri")}
    out = {}
    for row in aug["Nodes"]:
        name = s(row.get("Node name"))
        uri = None
        if name in votes:
            uri = max(votes[name].items(), key=lambda kv: kv[1])[0]
        if not uri:
            want = norm(name)
            hits = [u for n, u in by_name.items() if n.startswith(want[:25]) or want.startswith(n[:25])]
            uri = hits[0] if len(hits) == 1 else None
        if uri:
            out[uri] = row
    return out


def urls_in(v) -> list[str]:
    return [u.rstrip(".,;)") for u in re.findall(r"https?://[^\s<>\"']+", s(v))]
